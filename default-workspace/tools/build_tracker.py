# -*- coding: utf-8 -*-
"""生成《学习计划-进度追踪表.xlsx》初版。

注意：本脚本会整本重写工作簿，跑完后「周材料」工作表会丢。
改了周次或工时请依次执行：build_tracker.py → build_materials.py
"""
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

START = date(2026, 10, 19)
WEEKS = 38

NAVY = "1F3A5F"
BLUE = "2E5E8C"
LIGHT = "EAF1F8"
GREY = "F2F4F7"
GREEN = "C6EFCE"
YELLOW = "FFEB9C"
RED = "FFC7CE"

thin = Side(style="thin", color="D0D7DE")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

# (阶段, 主线任务, 周末交付, 计划工时)
DATA = [
    ("阶段 0", "Python 工程化：uv/pyproject、类型注解 + Pydantic、async/httpx；调通一家 LLM API 并实现流式输出",
     "CLI 流式问答：超时 / 重试 / --debug 打印 token 用量", 10),
    ("阶段 0", "FastAPI 路由与依赖注入；pytest + fixture + mock LLM；Docker 基础",
     "/chat 服务 + 8 条 pytest 全绿 + docker 一条命令起得来", 10),

    ("阶段 1", "文档加载与切分（PDF/MD，表格图片）、Embedding 原理、Chroma/FAISS 本地索引",
     "50 份 PDF 建索引，按引用回答问题", 10),
    ("阶段 1", "检索评估：先建 30 条评测集（问题 + 标准答案 + 应命中 chunk），再做 chunk 策略对比",
     "3 种 chunk 策略的 Recall@5 / MRR 对比表", 10),
    ("阶段 1", "Advanced RAG：Query 改写（HyDE/Multi-Query）、BM25 + 向量混合检索、Reranker",
     "混合检索 + rerank 后 Recall@5 相对基线提升 ≥10%", 10),
    ("阶段 1", "生成侧评估（RAGAs：Faithfulness / Answer Relevancy）、Milvus 单机部署、Docker 打包",
     "端到端系统 + 评估报告（基线 vs 优化版四指标）【M1】", 10),

    ("阶段 2", "不用框架，手写 Agent Loop（150 行内）：dispatch 表注册工具、最大步数、超时、结构化输出解析",
     "mini-agent：3 个工具 + 完整消息轨迹打印", 10),
    ("阶段 2", "Function Calling 深入、工具 schema 与描述设计、错误重试与降级策略",
     "工具失败率对比实验：无重试 / 指数退避 / 降级兜底", 10),
    ("阶段 2", "MCP：写一个自己的 MCP Server（本地文档检索），接入任意 MCP 客户端",
     "可被客户端调用的 Server + 契约测试", 10),
    ("阶段 2", "Memory（短期 / 会话 / 长期）+ Context Engineering：分层上下文、压缩触发、检索注入",
     "跨会话记住用户偏好的助手 + 压缩前后 token 对比", 10),
    ("阶段 2", "LangGraph 状态机：中断、人工确认、恢复；拆解一个真实 Harness 源码（Claude Code / OpenHands）",
     "research-agent（RAG + 搜索 + 引用）+ 一页 Harness 拆解笔记【M2】", 10),

    ("阶段 3", "性能剖析（py-spy / cProfile）、Redis 缓存（LLM 响应 + Embedding）、异步化改造",
     "优化前后 QPS / P99 对比表（目标 QPS ×3）", 10),
    ("阶段 3", "批处理、模型路由（贵/便宜分级）、Token 成本核算、vLLM 部署开源小模型",
     "成本报表：单请求成本从 X 降到 Y", 10),
    ("阶段 3", "可观测：Langfuse/LangSmith trace、OpenTelemetry、Prometheus 自定义指标",
     "Grafana 大盘 + 3 个自定义指标（token / 缓存命中 / 工具失败率）", 10),
    ("阶段 3", "Docker Compose 全栈编排、GitHub Actions（pytest + eval）、故障注入演练",
     "一键起全栈 + 故障演练报告（杀 Milvus / 限流 / 超时）【M3】", 10),

    ("弹性周", "春节弹性周：不排新内容。补欠账、复习前半程、回看两个仓库的 README",
     "复盘笔记：前半程最大的 3 个坑 + 后半程调整项", 10),

    ("阶段 4", "三层评测（组件 / 轨迹 / 端到端）、任务集构造、rubric 设计、LLM-as-judge 偏差与校准",
     "50 条评测集（含难度分层、期望结果、判定方式）", 10),
    ("阶段 4", "Eval Harness：并发跑批、结果落库、失败自动分类、重复 3 次测方差、CI 集成",
     "一条命令输出：通过率 / 步数 / 成本 / 延迟 / 失败分布【M4】", 10),

    ("阶段 5", "Multi-Agent 判断标准、Supervisor / Graph 编排、上下文隔离、循环与漂移治理",
     "planner → executor → reviewer 小系统 + 一页决策笔记", 10),
    ("阶段 5", "项目 1 开工：需求、架构、数据源、工具集；先写 20 条验收用例再写代码",
     "架构图 + 接口定义 + 20 条验收用例", 10),
    ("阶段 5", "项目 1 收尾：评测、监控、部署、README（写清取舍与限制）",
     "可运行 + 可评测 + 可部署", 10),
    ("阶段 5", "项目 2 攻坚：把 Eval Harness 产品化（Web 界面 + 轨迹回放 + 版本对比）",
     "能上传 case、跑批、看 diff 的界面【M5】", 10),

    ("阶段 6", "简历重写（三维模板）、项目陈述练两版（3 分钟 / 10 分钟）、系统设计 8 题",
     "简历 v3 + 项目讲解稿", 10),
    ("阶段 6", "模拟面试 2 场、高频题清单过一遍、开始投递 + 找内推",
     "投递记录表（岗位 / 进度 / 反馈）", 10),

    ("阶段 B", "具身全景：感知-决策-控制闭环、VLA 是什么、与 LLM Agent 的关系；PyTorch 补课",
     "一页 VLA 类比笔记 + 手写 MLP 训练脚本", 10),
    ("阶段 B", "机器人学最小集：坐标系与变换、正逆运动学、关节/任务空间、URDF；ROS 2 基础",
     "2 连杆逆解 + 一个 ros2 最小节点跑通", 10),
    ("阶段 B", "仿真入门：MuJoCo、Gymnasium 接口、奖励设计",
     "机械臂 reach 任务：随机策略 vs 训练策略曲线", 10),
    ("阶段 B", "强化学习最小集：PPO 原理与调参、stable-baselines3、域随机化概念",
     "reach 成功率 0 → 80%+，附训练曲线", 10),
    ("阶段 B", "模仿学习：行为克隆、ACT、Diffusion Policy；LeRobot 数据集格式与训练流程",
     "训一个抓取 policy + 回放视频", 10),
    ("阶段 B", "SO-101 + LeRobot 遥操作（无硬件用仿真替代）：标定、示教、数据采集",
     "50 条示教数据 + 训练 + 实机/仿真回放", 10),
    ("阶段 B", "VLA 论文精读 1：RT-2 / OpenVLA（VLM backbone + action token 化）",
     "论文笔记：输入输出、动作表示、训练数据、评测方式", 10),
    ("阶段 B", "VLA 论文精读 2：π0 / GR00T（flow matching、action expert、跨本体）",
     "两代架构对比表：动作表示 / 推理频率 / 数据规模", 10),
    ("阶段 B", "复现：LIBERO 基准上跑 OpenVLA 或 π0 的 zero-shot 评测",
     "复现报告：成功率 / 延迟 / 显存 / 与论文数字的差异", 10),
    ("阶段 B", "微调：LoRA 微调 VLA 到自定义任务（仿真数据即可）+ 数据量消融",
     "微调前后成功率对比 + 10/50/200 条消融", 10),
    ("阶段 B", "Isaac Sim / Isaac Lab 入门：USD 资产、场景搭建、GPU 并行环境",
     "Isaac Lab 里搭一个抓取场景并跑通训练", 10),
    ("阶段 B", "Sim2Real：域随机化、相机标定、延迟与噪声建模、真机部署（SO-101）",
     "仿真训练 → 真机 demo + 失败案例记录", 10),
    ("阶段 B", "具身评测：LIBERO/CALVIN 基准、成功判定标准、数据质量、可复现性",
     "评测报告：同一 policy 在 3 种条件下成功率差异与归因", 10),
    ("阶段 B", "收尾：技术报告、3 分钟演示视频、作品集整合、简历更新",
     "demo 视频 + 技术报告 + 更新版简历【M6】", 10),
]
assert len(DATA) == WEEKS, len(DATA)

wb = Workbook()

# ---------------- 总览 ----------------
ws = wb.active
ws.title = "总览"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 30
for c in "BCDEF":
    ws.column_dimensions[c].width = 14

ws.merge_cells("A1:F1")
ws["A1"] = "Agent / 具身智能 38 周学习计划 · 进度追踪"
ws["A1"].font = Font(size=16, bold=True, color="FFFFFF")
ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
ws["A1"].alignment = Alignment(horizontal="left", vertical="center", indent=1)
ws.row_dimensions[1].height = 38

ws.merge_cells("A2:F2")
ws["A2"] = "测试开发转型 ｜ 每周 10 小时 ｜ 2026-10-19 起 ｜ 预期 2027-07-11 完成"
ws["A2"].font = Font(size=10, color="44546A")
ws["A2"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

ws.merge_cells("A3:F3")
ws["A3"] = "用法：每周日填写「周计划」里的实际工时、状态、自评；本页自动汇总。状态只有 4 个值，别自创。"
ws["A3"].font = Font(size=10, italic=True, color="7F7F7F")
ws["A3"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

metrics = [
    ("计划总周数", "=COUNTA(周计划!A2:A39)", "0"),
    ("计划总工时", "=SUM(周计划!G2:G39)", "0"),
    ("已完成周数", '=COUNTIF(周计划!I2:I39,"已完成")', "0"),
    ("进行中周数", '=COUNTIF(周计划!I2:I39,"进行中")', "0"),
    ("延期周数", '=COUNTIF(周计划!I2:I39,"延期")', "0"),
    ("完成率", "=IFERROR(B7/B5,0)", "0.0%"),
    ("累计实际工时", "=SUM(周计划!H2:H39)", "0"),
    ("剩余工时（计划-实际）", "=B6-B11", "0"),
    ("平均自评（1-5）", "=IFERROR(AVERAGE(周计划!J2:J39),0)", "0.0"),
]
r = 5
for label, formula, fmt in metrics:
    ws.cell(row=r, column=1, value=label).font = Font(bold=True, color=NAVY)
    c = ws.cell(row=r, column=2, value=formula)
    c.number_format = fmt
    c.font = Font(bold=True, size=12)
    c.alignment = Alignment(horizontal="left")
    for col in (1, 2):
        ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=LIGHT)
        ws.cell(row=r, column=col).border = BORDER
    r += 1

# 阶段汇总
r = 16
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
ws.cell(row=r, column=1, value="阶段汇总").font = Font(size=12, bold=True, color="FFFFFF")
for col in range(1, 7):
    ws.cell(row=r, column=col).fill = PatternFill("solid", fgColor=BLUE)
ws.row_dimensions[r].height = 24

head = ["阶段", "", "周数", "计划工时", "已完成周", "完成率"]
r += 1
for i, h in enumerate(head, start=1):
    c = ws.cell(row=r, column=i, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="5B7FA6")
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.border = BORDER
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)

stages = [
    ("阶段 0", "地基 · 工程化与可测试服务"),
    ("阶段 1", "RAG · 检索增强与检索评估"),
    ("阶段 2", "Agent 核心 · Loop / Tool / Memory"),
    ("阶段 3", "工程化 · 缓存 / 异步 / 观测 / 部署"),
    ("弹性周", "弹性周 · 春节缓冲"),
    ("阶段 4", "评测 · 三层评测 + Eval Harness"),
    ("阶段 5", "Multi-Agent 与两个项目"),
    ("阶段 6", "求职冲刺 · 简历与投递"),
    ("阶段 B", "具身 · 仿真 → 模仿学习 → VLA"),
]
start = r + 1
for i, (st, name) in enumerate(stages):
    row = start + i
    name_cell = ws.cell(row=row, column=1, value=f"{st}　{name}")
    name_cell.alignment = Alignment(vertical="center", wrap_text=True, indent=1)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    ws.cell(row=row, column=3, value=f'=COUNTIF(周计划!$D$2:$D$39,"{st}")')
    ws.cell(row=row, column=4, value=f'=SUMIF(周计划!$D$2:$D$39,"{st}",周计划!$G$2:$G$39)')
    ws.cell(row=row, column=5, value=f'=COUNTIFS(周计划!$D$2:$D$39,"{st}",周计划!$I$2:$I$39,"已完成")')
    ws.cell(row=row, column=6, value=f"=IFERROR(E{row}/C{row},0)")
    ws.cell(row=row, column=6).number_format = "0%"
    ws.row_dimensions[row].height = 26
    for col in range(1, 7):
        cell = ws.cell(row=row, column=col)
        cell.border = BORDER
        if i % 2 == 1:
            cell.fill = PatternFill("solid", fgColor=GREY)
        if col >= 3:
            cell.alignment = Alignment(horizontal="center", vertical="center")

# ---------------- 周计划 ----------------
ws2 = wb.create_sheet("周计划")
ws2.sheet_view.showGridLines = False
headers = ["周次", "开始日期", "结束日期", "阶段", "主线任务（工作日 5×1h 输入）",
           "周末交付（可验收）", "计划工时", "实际工时", "状态", "自评(1-5)", "备注 / 卡点"]
widths = [7, 11.5, 11.5, 9.5, 54, 42, 9, 9, 9.5, 11.5, 28]
for i, (h, w) in enumerate(zip(headers, widths), start=1):
    c = ws2.cell(row=1, column=i, value=h)
    c.font = Font(bold=True, color="FFFFFF", size=10)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BORDER
    ws2.column_dimensions[get_column_letter(i)].width = w
ws2.row_dimensions[1].height = 34

for i, (stage, task, deliverable, hours) in enumerate(DATA):
    row = i + 2
    s = START + timedelta(weeks=i)
    e = s + timedelta(days=6)
    vals = [i + 1, s, e, stage, task, deliverable, hours, None, "未开始", None, None]
    for j, v in enumerate(vals, start=1):
        c = ws2.cell(row=row, column=j, value=v)
        c.border = BORDER
        c.font = Font(size=10)
        c.alignment = Alignment(vertical="center", wrap_text=j in (5, 6, 11),
                                horizontal="center" if j in (1, 4, 7, 8, 9, 10) else "left")
    ws2.cell(row=row, column=2).number_format = "yyyy-mm-dd"
    ws2.cell(row=row, column=3).number_format = "yyyy-mm-dd"
    ws2.row_dimensions[row].height = 46
    if i % 2 == 1:
        for j in range(1, 12):
            ws2.cell(row=row, column=j).fill = PatternFill("solid", fgColor=GREY)

dv = DataValidation(type="list", formula1='"未开始,进行中,已完成,延期"', allow_blank=False,
                    showDropDown=False, error="状态只能是：未开始 / 进行中 / 已完成 / 延期",
                    errorTitle="状态不合法")
ws2.add_data_validation(dv)
dv.add("I2:I39")

dv2 = DataValidation(type="whole", operator="between", formula1=1, formula2=5, allow_blank=True,
                     error="自评填 1-5 的整数", errorTitle="超出范围")
ws2.add_data_validation(dv2)
dv2.add("J2:J39")

ws2.conditional_formatting.add("I2:I39", CellIsRule(operator="equal", formula=['"已完成"'],
                                                     fill=PatternFill("solid", fgColor=GREEN)))
ws2.conditional_formatting.add("I2:I39", CellIsRule(operator="equal", formula=['"进行中"'],
                                                     fill=PatternFill("solid", fgColor=YELLOW)))
ws2.conditional_formatting.add("I2:I39", CellIsRule(operator="equal", formula=['"延期"'],
                                                     fill=PatternFill("solid", fgColor=RED)))
ws2.freeze_panes = "A2"
ws2.auto_filter.ref = "A1:K39"

# ---------------- 里程碑 ----------------
ws3 = wb.create_sheet("里程碑")
ws3.sheet_view.showGridLines = False
ws3.merge_cells("A1:F1")
ws3["A1"] = "六个里程碑 · 验收清单"
ws3["A1"].font = Font(size=14, bold=True, color="FFFFFF")
ws3["A1"].fill = PatternFill("solid", fgColor=NAVY)
ws3["A1"].alignment = Alignment(vertical="center", indent=1)
ws3.row_dimensions[1].height = 32

m_head = ["里程碑", "目标日期", "对应周", "验收标准", "状态", "备注"]
m_widths = [22, 13, 9, 58, 10, 26]
for i, (h, w) in enumerate(zip(m_head, m_widths), start=1):
    c = ws3.cell(row=3, column=i, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=BLUE)
    c.alignment = Alignment(horizontal="center")
    c.border = BORDER
    ws3.column_dimensions[get_column_letter(i)].width = w

MILES = [
    ("M1　RAG 系统", date(2026, 11, 29), "W1-W6", "30 条评测集可复跑，优化前后指标有对比表"),
    ("M2　手写 Agent", date(2027, 1, 3), "W7-W11", "150 行内跑通 loop，工具 / 重试 / trace 齐全"),
    ("M3　可部署系统", date(2027, 1, 31), "W12-W15", "compose up 一键起，有 QPS / P99 / 成本三项数据"),
    ("M4　Eval Harness", date(2027, 2, 21), "W17-W18", "一条命令输出通过率 / 步数 / 成本 / 失败分布"),
    ("M5　两个简历项目", date(2027, 3, 21), "W19-W22", "可运行 + 可评测 + README 有取舍说明"),
    ("M6　具身 pipeline", date(2027, 7, 11), "W25-W38", "仿真复现 + 微调 + sim2real + 评测报告 + 视频"),
]
for i, (name, d, wk, std) in enumerate(MILES):
    row = 4 + i
    for j, v in enumerate([name, d, wk, std, "未开始", None], start=1):
        c = ws3.cell(row=row, column=j, value=v)
        c.border = BORDER
        c.alignment = Alignment(vertical="center", wrap_text=j in (1, 4, 6),
                                horizontal="center" if j in (2, 3, 5) else "left")
        if i % 2 == 1:
            c.fill = PatternFill("solid", fgColor=GREY)
    ws3.cell(row=row, column=2).number_format = "yyyy-mm-dd"
    ws3.row_dimensions[row].height = 30

dv3 = DataValidation(type="list", formula1='"未开始,进行中,已达成,未达成"', allow_blank=False,
                     showDropDown=False)
ws3.add_data_validation(dv3)
dv3.add("E4:E9")

# ---------------- 资源清单 ----------------
ws4 = wb.create_sheet("资源清单")
ws4.sheet_view.showGridLines = False
ws4.merge_cells("A1:E1")
ws4["A1"] = "资源清单（每个阶段只留 3 个，贪多是最大的坑）"
ws4["A1"].font = Font(size=14, bold=True, color="FFFFFF")
ws4["A1"].fill = PatternFill("solid", fgColor=NAVY)
ws4["A1"].alignment = Alignment(vertical="center", indent=1)
ws4.row_dimensions[1].height = 32

r_head = ["阶段", "资源", "链接", "用途", "优先级"]
r_widths = [10, 34, 62, 40, 8]
for i, (h, w) in enumerate(zip(r_head, r_widths), start=1):
    c = ws4.cell(row=3, column=i, value=h)
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor=BLUE)
    c.alignment = Alignment(horizontal="center")
    c.border = BORDER
    ws4.column_dimensions[get_column_letter(i)].width = w

RES = [
    ("阶段 0-2", "Anthropic: Building effective agents", "https://www.anthropic.com/engineering/building-effective-agents", "Agent 与 Workflow 的边界，必读三遍", "P0"),
    ("阶段 0-2", "OpenAI: A practical guide to building agents", "https://openai.com/business/guides-and-resources/a-practical-guide-to-building-ai-agents/", "工程视角的 Agent 设计", "P1"),
    ("阶段 0-2", "Hello-Agents（Datawhale）", "https://datawhalechina.github.io/hello-agents/", "中文 Agent 主线教材", "P0"),
    ("阶段 1", "All-in-RAG（Datawhale）", "https://github.com/datawhalechina/all-in-rag", "RAG 系统化实现", "P0"),
    ("阶段 1", "RAG from Scratch", "https://github.com/langchain-ai/rag-from-scratch", "检索技术速查", "P1"),
    ("阶段 1", "RAGAs", "https://docs.ragas.io/", "RAG 自动评估", "P0"),
    ("阶段 2", "MCP 协议", "https://modelcontextprotocol.io/", "工具生态事实标准", "P0"),
    ("阶段 2", "LangGraph", "https://langchain-ai.github.io/langgraph/", "状态机、中断、人工确认", "P0"),
    ("阶段 2", "learn-claude-code", "https://github.com/shareAI-lab/learn-claude-code", "从零复刻 coding agent harness", "P1"),
    ("阶段 3", "Langfuse", "https://langfuse.com/", "开源 trace，适合自建", "P0"),
    ("阶段 3", "vLLM", "https://docs.vllm.ai/", "推理加速与吞吐测试", "P1"),
    ("阶段 4", "DeepEval", "https://github.com/confident-ai/deepeval", "断言式评测", "P1"),
    ("阶段 B", "Foundation_of_Embodied_AI", "https://github.com/zengxiong111/Foundation_of_Embodied_AI", "中文：机器人学 + Isaac + MuJoCo + RL", "P0"),
    ("阶段 B", "awesome-vla-study", "https://github.com/MilkClouds/awesome-vla-study", "VLA 论文阅读顺序", "P0"),
    ("阶段 B", "LeRobot（HuggingFace）", "https://github.com/huggingface/lerobot", "数据格式 + 训练 + SO-101", "P0"),
    ("阶段 B", "openpi（π0）", "https://github.com/Physical-Intelligence/openpi", "开源 VLA 代表之一", "P0"),
    ("阶段 B", "OpenVLA", "https://github.com/openvla/openvla", "开源 VLA 代表之二", "P1"),
    ("阶段 B", "LIBERO", "https://github.com/Lifelong-Robot-Learning/LIBERO", "语言条件操作基准，复现首选", "P0"),
    ("阶段 B", "MuJoCo + Gymnasium + SB3", "https://github.com/google-deepmind/mujoco", "RL 最小闭环", "P0"),
    ("阶段 B", "Isaac Lab", "https://github.com/isaac-sim/IsaacLab", "GPU 并行仿真与 sim2real", "P1"),
    ("阶段 B", "NVIDIA Physical AI 课程", "https://docs.nvidia.com/learning/physical-ai/robotics.html", "含 SO-101 sim2real 完整流程", "P1"),
]
for i, row_vals in enumerate(RES):
    row = 4 + i
    stage, name, url, use, prio = row_vals
    display = url if len(url) <= 72 else url.split("//")[1].split("/")[0] + "/…"
    for j, v in enumerate([stage, name, display, use, prio], start=1):
        c = ws4.cell(row=row, column=j, value=v)
        c.border = BORDER
        c.font = Font(size=10)
        c.alignment = Alignment(vertical="center", wrap_text=j in (2, 3, 4),
                                horizontal="center" if j in (1, 5) else "left")
    link = ws4.cell(row=row, column=3)
    link.hyperlink = url
    link.font = Font(size=10, color="0563C1", underline="single")
    if i % 2 == 1:
        for j in range(1, 6):
            ws4.cell(row=row, column=j).fill = PatternFill("solid", fgColor=GREY)
    ws4.row_dimensions[row].height = 28

wb.save("学习计划-进度追踪表.xlsx")
print("saved")
