# -*- coding: utf-8 -*-
"""生成《成果验收标准-38周.md》，并把「成果验收」工作表写入进度追踪表。

数据是唯一来源：改 WEEKS 后重跑，md 与 xlsx 一起更新。
"""

MD_OUT = "成果验收标准-38周.md"
XLSX = "学习计划-进度追踪表.xlsx"

# (周, 日期, 一句话成果, 交付证据, 达标线, 不合格的样子)
WEEKS = [
    (1, "10/19–10/25",
     "能用工程化方式调用 LLM，并说清 token 计费与首 token 延迟",
     "mini-llm-service 仓库首个 commit；`ask` 命令的录屏或 gif；一页笔记（流式与 SSE 的关系）",
     "`--stream` 逐字输出；`--debug` 打印 prompt/completion 两个 token 数；断网或超时 3 秒内报错且不崩",
     "直接 print 整段返回；没有超时配置；token 数靠估"),

    (2, "10/26–11/01",
     "能交付一个「别人 clone 下来就能跑、能测」的服务",
     "/chat 服务 + 8 条 pytest + Dockerfile；README 里两条命令（build / run）",
     "`pytest -q` 全绿且不需要真 key（全 mock）；镜像 < 1.2GB；curl 拿到流式",
     "测试要真 key、跑一次花两毛；Dockerfile 用 latest；没有 .dockerignore"),

    (3, "11/02–11/08",
     "能把 50 份 PDF 变成可检索、可引用答案的知识库",
     "索引构建脚本 + 5 问 5 答的引用截图 + 解析失败清单",
     "抽 3 个含表格的 PDF 人工核对不丢内容；每条答案附 chunk 原文；重建索引 < 5 分钟",
     "只支持 txt；引用写「来自文档」这种废话；表格解析成一团"),

    (4, "11/09–11/15",
     "有一把尺子，能证明检索变好还是变坏",
     "30 条评测集（JSONL）+ 评测脚本 + 三种策略对比表",
     "人工确认 ≥20 条有效；一条命令跑完；表里同时有 Recall@5 与 MRR",
     "让模型自己出题自己答；只写「感觉更准了」"),

    (5, "11/16–11/22",
     "能用改写、混合检索、重排把指标提上去，并说清各自代价",
     "四行消融表（基线 / +改写 / +混合 / +重排）+ 各方案延迟",
     "Recall@5 相对基线 ≥ +10%；每个改动的 P99 延迟增量有记录",
     "全开但不做消融，说不出哪一步真的有用"),

    (6, "11/23–11/29",
     "能证明「答案没瞎编」，并把系统打包交付",
     "RAGAs 四指标报告（基线 vs 优化）+ Milvus 部署脚本 + v0.1 tag",
     "Faithfulness ≥ 0.85（达不到就写清基线并说明原因）；三条命令可复现",
     "只贴一个分数，没有对照组"),

    (7, "11/30–12/06",
     "不看框架写出 150 行 Agent Loop，并能解释循环为什么停",
     "mini-agent 仓库 + 一次完整轨迹（JSONL，含每步耗时）",
     "3 个工具、最大步数、超时、解析失败重试齐全；轨迹能逐步回放",
     "`while True` 没有上限；工具调用写死在 prompt 描述里"),

    (8, "12/07–12/13",
     "工具 30% 失败率下，Agent 仍能完成任务",
     "失败注入开关 + 三策略对比表（无重试 / 退避 / 降级）",
     "成功率 ≤70% → ≥95%；平均步数增幅 < 30%；成本不翻倍",
     "只包了 try/except，没有降级路径"),

    (9, "12/14–12/20",
     "能力能以 MCP Server 形式被任意客户端复用",
     "Server 仓库 + 契约测试 + 客户端调用截图",
     "5 条契约用例通过（含非法参数）；错误信息可读；README 有 3 条示例",
     "只能在自己脚本里 import，谈不上协议"),

    (10, "12/21–12/27",
     "Agent 跨会话记得用户，且上下文不炸",
     "记忆库 + 20 轮对话日志 + 压缩前后 token 对比",
     "第 2 天能召回偏好；20 轮后 token 增长 < 2 倍；10 条抽查召回准确率 ≥80%",
     "把所有历史一股脑塞进 prompt"),

    (11, "12/28–01/03",
     "能让 Agent 停下来等人确认，也能读懂别人的 Harness",
     "research-agent 仓库 + 中断演示 + 一页拆解笔记",
     "高风险动作必须人工确认；进程被杀后能续跑；笔记指出 4 个机制的具体位置",
     "「多 Agent 就是多写几个 prompt」"),

    (12, "01/04–01/10",
     "拿得出「优化前后」的硬指标",
     "火焰图 + QPS/P99 前后对比 + 改动清单",
     "QPS ≥3 倍或 P99 降 ≥50%，且成功率不变；每个改动能单独开关",
     "换成更贵的模型充数，还说「性能提升了」"),

    (13, "01/11–01/17",
     "能把单次成本降下来，并说清代价",
     "成本拆解表（精确到分）+ 路由策略 + 质量对比",
     "单请求成本降 ≥40%；任务成功率下降 < 3%",
     "只关缓存不测质量"),

    (14, "01/18–01/24",
     "线上出问题，5 分钟内定位到具体哪一步",
     "trace 截图 + Grafana 大盘 + 2 条告警规则",
     "注入一次超时，仅凭 trace 定位；3 个业务指标 + 2 条告警可触发",
     "只有 print 日志，全链路靠猜"),

    (15, "01/25–01/31",
     "一键起全栈，且扛住一次人为故障",
     "compose 文件 + CI 截图 + 故障演练报告",
     "`compose up` 90 秒内全绿；CI 5 分钟内跑完 pytest + eval；4 类故障有预期行为",
     "手工起服务；故障靠重启解决"),

    (16, "02/01–02/07",
     "知道自己到哪了，欠账清零",
     "复盘笔记 + 指标总表（QPS / P99 / 成本 / 成功率）",
     "四个核心数字都有出处；所有仓库 README 都能跑",
     "用这周追进度，结果新债叠旧债"),

    (17, "02/08–02/14",
     "有一套可判定的评测集，而不是「感觉」",
     "50 条任务集 + rubric 文档 + judge 校准记录",
     "5 类任务、4 档 rubric；人工标 30 条，judge 与人不一致率 < 15%",
     "只用「答案是否含关键词」判定"),

    (18, "02/15–02/21",
     "一条命令跑完回归，输出可对比的报告",
     "eval-harness v0.1 + 一次真实运行报告",
     "100 条 ≤3 分钟；输出通过率/步数/成本/延迟/失败分布；重复 3 次方差 < 5%",
     "每次手动跑 notebook，结果存桌面"),

    (19, "02/22–02/28",
     "能判断「这里不该用多 Agent」，并给出替代方案",
     "3 角色小系统 + 决策笔记（4 条判据 + 2 个反例）",
     "3 个角色的输入输出 schema 固定；反例能说明退回单 Agent 的收益",
     "为了简历堆角色，效果不如单 Agent"),

    (20, "03/01–03/07",
     "写代码之前，手里已有 20 条验收用例",
     "设计文档 + 20 条用例（此刻全部失败）",
     "每条有输入、期望、判定方式；跑一遍能看到 20 条红",
     "先把功能写完，再补「用例」"),

    (21, "03/08–03/14",
     "项目 1 的通过率曲线是真的",
     "20 条用例过 18+ + 四项指标表 + 部署脚本",
     "成功率 ≥90%；指标四项齐全；失败能归到类别",
     "挑简单用例凑数"),

    (22, "03/15–03/21",
     "不懂命令行的人也能跑一次回归",
     "Web 界面 + 轨迹回放 + 版本 diff + 3 分钟视频",
     "上传 YAML → 跑批 → 看 diff 全流程 ≤5 分钟；报告可导出",
     "只有 CLI，还要求用户会装 Python"),

    (23, "03/22–03/28",
     "项目能被 3 分钟讲清楚，且经得住追问",
     "简历 v3 + 3 分钟 / 10 分钟两版讲稿 + 录音",
     "不看稿讲完；每个数字说得出出处；准备好 3 个「为什么不」的答案",
     "简历写「熟悉 LangChain」"),

    (24, "03/29–04/04",
     "第一批投递已发出，反馈已记录",
     "投递记录表（≥5 家）+ 2 场 mock 录像",
     "每家都有反馈字段；每场 mock 收 3 条改进项",
     "「再准备准备就投」"),

    (25, "04/05–04/11",
     "能训一个模型，并讲清 VLA 与 LLM Agent 的差别",
     "MLP 训练脚本 + 一页笔记 + 3 分钟口述录音",
     "loss 收敛；能说出「动作不是文本」的 3 个理由",
     "只跑通官方 demo"),

    (26, "04/12–04/18",
     "能算出关节角，并让 ROS 2 节点动起来",
     "逆解脚本 + rclpy 节点 + RViz 截图",
     "逆解误差 < 1e-3；节点能发布关节指令；URDF 可视化正确",
     "只跑 turtlesim"),

    (27, "04/19–04/25",
     "有自己的仿真环境，且基线是 0",
     "环境代码 + 随机策略基线曲线",
     "1000 步不崩；接口符合 Gymnasium 规范；成功率接近 0（此时正常）",
     "拿现成环境改个名字"),

    (28, "04/26–05/02",
     "把 reach 训到 80%，并知道为什么",
     "训练曲线 + 成功率 + 3 条失败分析",
     "成功率 ≥80%；同种子复现波动 ±5% 内",
     "只贴一张好看的上升曲线"),

    (29, "05/03–05/09",
     "能训一个抓取 policy，并解释行为克隆为什么会失效",
     "训练脚本 + 回放视频 + 失效分析",
     "视频里至少 1 次成功抓取；写出 3 个失效原因（复合误差、数据偏置、分布漂移）",
     "训练 loss 下降就宣布成功"),

    (30, "05/10–05/16",
     "有真机（或仿真等价物）的数据闭环",
     "50 条示教数据 + 训练脚本 + 实机回放视频",
     "数据字段完整、能被 LeRobot 加载；实机至少 1 次成功",
     "数据是拼凑的，无法复现"),

    (31, "05/17–05/23",
     "能讲清 OpenVLA 的四个问题：输入、输出、数据、评测",
     "论文笔记 + 5 分钟讲解录音",
     "四问都能答；能指出 2 个可疑的实验设置",
     "复述摘要"),

    (32, "05/24–05/30",
     "能解释 flow matching 为什么取代离散 action token",
     "架构对比表（4 列齐全）+ 一页解释",
     "说清推理频率与 action chunk 的关系；说清两代的数据规模差异",
     "只会说「π0 更强」"),

    (33, "05/31–06/06",
     "复现了 zero-shot，并知道差在哪",
     "复现报告（我的数字 vs 论文数字）",
     "跑通 10 个任务；差异有归因；别人照文档能复现",
     "只截图成功的那一次"),

    (34, "06/07–06/13",
     "能用小数据把模型微调出可测的提升",
     "微调脚本 + 三档数据量消融曲线 + 对比表",
     "微调后成功率提升 ≥10 个百分点；10/50/200 三个点齐全",
     "全量微调烧钱，不做消融"),

    (35, "06/14–06/20",
     "能在 Isaac 里搭场景并并行训练",
     "场景文件 + 训练日志 + 回放",
     "64 环境并行跑通；显存与耗时记录在案",
     "只打开了 GUI 截图"),

    (36, "06/21–06/27",
     "仿真训出来的 policy 在真机上动了",
     "sim2real 视频 + 标定参数 + 失败清单",
     "至少 1 个任务端到端跑通；5 条失败案例各有假设",
     "只在仿真里「看起来能行」"),

    (37, "06/28–07/04",
     "能给出一份可信的成功率，而不是一个数字",
     "评测协议 + 3 组条件对比 + 方差或置信区间",
     "每组 ≥20 次试验；给出波动范围；协议别人能照做",
     "跑 3 次取平均"),

    (38, "07/05–07/11",
     "38 周被压缩成 3 分钟",
     "demo 视频 + 技术报告 + 作品集页 + 简历 v4",
     "外行看 3 分钟能说出你做了什么；3 个仓库置顶可点",
     "只剩一堆散乱的 notebook"),
]


def build_md():
    L = []
    L.append("# 成果验收标准（38 周）")
    L.append("")
    L.append("> 配套：[学习计划-Agent与具身智能-38周.md](学习计划-Agent与具身智能-38周.md)｜[每周学习材料-38周.md](每周学习材料-38周.md)｜[学习计划-进度追踪表.xlsx](学习计划-进度追踪表.xlsx)")
    L.append("> 这份文档只回答一个问题：**做到什么程度才算这一周没白过。**")
    L.append("")
    L.append("## 一、先把「成果」定义清楚")
    L.append("")
    L.append("成果不是「学了什么」，是**别人能验证的东西**。只有四种形态：")
    L.append("")
    L.append("| 形态 | 例子 | 怎么验证 |")
    L.append("|:--|:--|:--|")
    L.append("| 能跑 | 一条命令、一个接口、一个脚本 | 别人 clone 下来执行不报错 |")
    L.append("| 能看 | 一张曲线图、一张对比表、一张 trace 截图 | 数据有出处，不是手画的 |")
    L.append("| 能读 | README、评测报告、设计文档 | 三个月后的你看得懂 |")
    L.append("| 能讲 | 3 分钟讲清取舍 | 不看稿、不被问倒 |")
    L.append("")
    L.append("**关键判据**：任何一个数字，你都要能回答「这个数是怎么来的」。面试官只问这一句。")
    L.append("")
    L.append("## 二、每天：一个成果 + 两个强制动作")
    L.append("")
    L.append("### 每天必须有 1 个成果（三选一）")
    L.append("")
    L.append("| 档位 | 标准 | 用时 | 常见样子 |")
    L.append("|:--|:--|:--|:--|")
    L.append("| 及格 | 读完 + 跑了最小例子 + 3 行笔记 | 1h | 一个能 print 出结果的脚本 |")
    L.append("| 良好 | 产物进了仓库，别人能照着跑 | 1.5h | 带 README 的小脚本 + 输出样例 |")
    L.append("| 优秀 | 顺手做了个「换成 X 会怎样」的小实验 | 2h | 两行对比数字，哪怕结论是「没差别」 |")
    L.append("")
    L.append("### 两个强制动作（不做就算没学）")
    L.append("")
    L.append("1. **一次 commit**：代码或笔记，进仓库。留在 notebook 里等于不存在。")
    L.append("2. **三行日志**：`做了什么` / `卡在哪` / `明天第一步`。写不满三行，说明今天没推进。")
    L.append("")
    L.append("### 一天的收尾（10 分钟）")
    L.append("")
    L.append("- 今天的产出能不能被截图放进周报？不能 → 明天补一个能跑的最小版本。")
    L.append("- 今天的「卡点」是知识缺口还是环境问题？环境问题当天解决，别带过夜。")
    L.append("")
    L.append("## 三、每周：六件套 + 一张验收单")
    L.append("")
    L.append("每周日填完这六项，才算交付：")
    L.append("")
    L.append("| # | 件 | 判定 |")
    L.append("|:--|:--|:--|")
    L.append("| 1 | 代码 | 仓库有 commit，且能被别人跑起来 |")
    L.append("| 2 | 数字 | 至少一个指标（指标 / 耗时 / 成本 / 成功率） |")
    L.append("| 3 | 文档 | README 或笔记更新，含「已知限制」 |")
    L.append("| 4 | 图 | 一张能直接放进 PPT 的图 |")
    L.append("| 5 | 话 | 3 分钟讲清本周做了什么、为什么 |")
    L.append("| 6 | 复盘 | 5 行周报 + 自评 1–5 |")
    L.append("")
    L.append("## 四、38 周成果清单")
    L.append("")
    L.append("每周一行：一句话成果 → 证据 → 达标线 → 不合格的样子。")
    L.append("")

    for w, dates, cap, ev, bar, bad in WEEKS:
        L.append(f"### W{w}｜{dates}")
        L.append("")
        L.append(f"**成果**：{cap}")
        L.append("")
        L.append(f"- 证据：{ev}")
        L.append(f"- 达标：{bar}")
        L.append(f"- 不合格：{bad}")
        L.append("")

    L.append("---")
    L.append("")
    L.append("## 五、里程碑复核（每到一个 M，回来看一次）")
    L.append("")
    L.append("| 里程碑 | 时间 | 证据包（缺一不可） |")
    L.append("|:--|:--|:--|")
    L.append("| M1 RAG | W6 (11/29) | 30 条评测集 + 指标对比表 + 可复跑脚本 |")
    L.append("| M2 手写 Agent | W11 (01/03) | 150 行 loop 仓库 + 轨迹 JSONL + 拆解笔记 |")
    L.append("| M3 可部署 | W15 (01/31) | compose 一键起 + QPS/P99/成本 + 故障报告 |")
    L.append("| M4 Eval Harness | W18 (02/21) | v0.1 仓库 + 一次真实运行报告 + CI 截图 |")
    L.append("| M5 两个项目 | W22 (03/21) | 两个仓库 + 两个视频 + 一篇技术文章 |")
    L.append("| M6 具身 | W38 (07/11) | 复现报告 + sim2real 视频 + 评测报告 |")
    L.append("")
    L.append("## 六、自评 1–5 分怎么打（别乱打分）")
    L.append("")
    L.append("| 分 | 含义 | 判断依据 |")
    L.append("|:--|:--|:--|")
    L.append("| 1 | 看了，没动手 | 没有 commit |")
    L.append("| 2 | 抄通了，但不知道为什么 | 能跑，但改一个参数就崩 |")
    L.append("| 3 | 自己写出来了 | 不看教程能复现，能解释每行 |")
    L.append("| 4 | 能评测、能做消融 | 有对照组，说得出哪一步有用 |")
    L.append("| 5 | 能讲清取舍，能给别人讲明白 | 能回答「为什么不用另一种方案」 |")
    L.append("")
    L.append("> 连续两周自评 ≤2 分：不是你不努力，是这一阶段的任务颗粒度不对——砍掉一个非核心项，重排。")
    L.append("")
    L.append("## 七、不达标怎么办（降级规则）")
    L.append("")
    L.append("| 情况 | 动作 |")
    L.append("|:--|:--|")
    L.append("| 某周只完成 3/6 件 | 保留代码 + 数字，砍掉文档和图，下周补 |")
    L.append("| 达标线差一点（如 Recall +8% 而非 +10%） | 如实记录，写清原因，继续走——不要为了凑数调参 |")
    L.append("| 整周没动 | 不补课，直接进下一周；把欠的那件写进「已知限制」 |")
    L.append("| 连续两周延期 | 砍掉本阶段的加分项（视频、博客），保主线 |")
    L.append("")
    L.append("## 八、这份东西最终怎么用")
    L.append("")
    L.append("面试时的标准句式：")
    L.append("")
    L.append("> 「问题是什么 → 我怎么设计 → 怎么验证有效 → 哪一步失败了、为什么」。")
    L.append("")
    L.append("你能把任意一周的「证据 + 达标线」讲成这四句，这一周就真的变成你的了。")
    L.append("")
    return "\n".join(L)


def write_xlsx():
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import CellIsRule

    wb = load_workbook(XLSX)
    if "成果验收" in wb.sheetnames:
        del wb["成果验收"]
    ws = wb.create_sheet("成果验收", 3)
    ws.sheet_view.showGridLines = False

    navy, blue, grey = "1F3A5F", "2E5E8C", "F2F4F7"
    green, yellow, red = "C6EFCE", "FFEB9C", "FFC7CE"
    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.merge_cells("A1:G1")
    ws["A1"] = "38 周成果验收：一句话成果 / 证据 / 达标线 / 不合格的样子"
    ws["A1"].font = Font(size=14, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=navy)
    ws["A1"].alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 32

    heads = ["周次", "日期", "本周成果（一句话）", "交付证据", "达标线", "不合格的样子", "状态"]
    widths = [7, 13, 34, 42, 44, 34, 9]
    for i, (h, w) in enumerate(zip(heads, widths), start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor=blue)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = border
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[3].height = 26

    for i, (w, dates, cap, ev, bar, bad) in enumerate(WEEKS):
        row = 4 + i
        for j, v in enumerate([w, dates, cap, ev, bar, bad, "未开始"], start=1):
            c = ws.cell(row=row, column=j, value=v)
            c.border = border
            c.font = Font(size=9)
            c.alignment = Alignment(vertical="top", wrap_text=j in (2, 3, 4, 5, 6),
                                    horizontal="center" if j in (1, 7) else "left")
        ws.row_dimensions[row].height = 58
        if i % 2 == 1:
            for j in range(1, 8):
                ws.cell(row=row, column=j).fill = PatternFill("solid", fgColor=grey)

    dv = DataValidation(type="list", formula1='"未开始,进行中,已达标,未达标"', allow_blank=False,
                        showDropDown=False, error="只能是：未开始 / 进行中 / 已达标 / 未达标",
                        errorTitle="状态不合法")
    ws.add_data_validation(dv)
    dv.add("G4:G41")

    for t, f in (("已达标", green), ("进行中", yellow), ("未达标", red)):
        ws.conditional_formatting.add("G4:G41",
                                      CellIsRule(operator="equal", formula=[f'"{t}"'],
                                                 fill=PatternFill("solid", fgColor=f)))

    ws.freeze_panes = "A4"
    ws.auto_filter.ref = "A3:G41"
    wb.save(XLSX)


if __name__ == "__main__":
    md = build_md()
    with open(MD_OUT, "w", encoding="utf-8") as f:
        f.write(md)
    write_xlsx()
    print("ok", md.count("\n"), len(WEEKS))
